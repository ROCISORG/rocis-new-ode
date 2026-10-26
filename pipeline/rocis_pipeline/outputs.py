"""Everything the pipeline writes: explorer CSVs, website data, reports."""
import csv
import json
from datetime import datetime
from pathlib import Path

from .model import STEP, Cohort

TIME_FORMAT = "%Y-%m-%d %H:%M:%S"


def write_explorer_csvs(cohort: Cohort, folder: Path) -> list:
    """Write large_N.csv and small_N.csv in the format the explorer has always used."""
    folder.mkdir(parents=True, exist_ok=True)
    stamps = set()
    for m in cohort.monitors:
        if m.total:
            current, last = min(m.total), max(m.total)
            while current <= last:
                stamps.add(current)
                current += STEP
    stamps = sorted(stamps)
    written = []
    for kind, attr in (("small", "total"), ("large", "large")):
        path = folder / f"{kind}_{cohort.number}.csv"
        with open(path, "w", encoding="utf-8", newline="") as file:
            writer = csv.writer(file)
            writer.writerow(["Date/Time"] + [m.column for m in cohort.monitors])
            for stamp in stamps:
                row = [stamp.strftime(TIME_FORMAT)]
                for m in cohort.monitors:
                    value = getattr(m, attr).get(stamp)
                    row.append("" if value is None else repr(value))
                writer.writerow(row)
        written.append(path)
    return written


def _round(value):
    return None if value is None else round(value, 1)


def write_site_cohort(cohort: Cohort, meta: dict, neighborhoods: dict, folder: Path) -> dict:
    """Write data/cohort-N.js on a regular 15-minute grid and return its manifest entry."""
    first, last = cohort.span()
    count = int((last - first) / STEP) + 1
    grid = [first + i * STEP for i in range(count)]
    monitors = []
    for m in cohort.monitors:
        place = neighborhoods.get(m.code) or neighborhoods.get(m.code.split("_")[0]) or {}
        monitors.append({
            "column": m.column,
            "hashId": m.hash_id,
            "code": m.code,
            "name": (place.get("name") or m.code).split(" / ")[0],
            "county": place.get("county", ""),
            "total": [_round(m.total.get(s)) for s in grid],
            "large": [_round(m.large.get(s)) for s in grid],
        })
    payload = {
        "cohort": cohort.number,
        "label": meta.get("label_date", ""),
        "start": first.strftime(TIME_FORMAT),
        "stepMinutes": 15,
        "points": count,
        "source": cohort.source,
        "monitors": monitors,
        "issues": [i.as_dict() for i in cohort.issues if i.public],
    }
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"cohort-{cohort.number}.js"
    text = "ROCIS.register(" + json.dumps(payload, separators=(",", ":")) + ");\n"
    path.write_text(text, encoding="utf-8")
    return {
        "cohort": cohort.number,
        "label": meta.get("label_date", ""),
        "first": first.strftime(TIME_FORMAT),
        "last": last.strftime(TIME_FORMAT),
        "monitors": len(monitors),
        "warnings": sum(i.level in ("warning", "error") for i in cohort.issues if i.public),
        "source": cohort.source,
        "file": path.name,
        "kb": round(len(text) / 1024),
    }


def write_manifest(entries: list, neighborhoods: dict, folder: Path) -> Path:
    entries = sorted(entries, key=lambda e: e["cohort"])
    payload = {"generated": datetime.now().strftime("%Y-%m-%d %H:%M"), "cohorts": entries,
               "neighborhoods": {c: {"name": n["name"], "county": n["county"]} for c, n in neighborhoods.items()}}
    path = folder / "cohorts.js"
    path.write_text("ROCIS.manifest(" + json.dumps(payload, separators=(",", ":")) + ");\n", encoding="utf-8")
    return path


def write_hash_report(cohort_number: int, rows: list, path: Path) -> Path:
    """The team's per-cohort format: 'Here are the hash IDs for Cohort N:' and a two-column table."""
    width = max([7] + [len(r["site_id"]) for r in rows])
    lines = [f"Here are the hash IDs for Cohort {cohort_number}:", "",
             f"| {'Site ID'.ljust(width)} | Hash ID |"]
    for r in rows:
        lines.append(f"| {r['site_id'].ljust(width)} | {r['hash_id'].rjust(7)} |")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def write_summary_workbook(cohort_number: int, summaries: list, issues: list, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = Workbook()
    sheet = book.active
    sheet.title = f"Cohort {cohort_number}"
    headers = ["Site ID", "Column", "Files", "Readings", "Readings in cohort dates", "First reading",
               "Last reading", "Lines dropped", "Repeated readings", "Median total", "90th percentile total",
               "Highest total", "15-minute values present"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for s in summaries:
        sheet.append([s.site_id, s.column, s.files, s.readings, s.kept, s.first, s.last, s.dropped_lines,
                      s.repeated, s.median_total, s.p90_total, s.max_total, round(s.coverage, 3)])
    for row in sheet.iter_rows(min_row=2):
        row[5].number_format = row[6].number_format = "yyyy-mm-dd hh:mm"
        row[12].number_format = "0%"
    for column, width in zip("ABCDEFGHIJKLM", [9, 14, 7, 10, 14, 17, 17, 9, 10, 11, 12, 11, 12]):
        sheet.column_dimensions[column].width = width
    notes = book.create_sheet("Checks")
    notes.append(["Level", "Message"])
    for cell in notes[1]:
        cell.font = Font(bold=True)
    for issue in issues:
        notes.append([issue.level, issue.message])
    notes.column_dimensions["A"].width = 10
    notes.column_dimensions["B"].width = 110
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def write_checks_report(cohorts: list, path: Path) -> Path:
    lines = ["# Data checks", "", f"Generated {datetime.now():%Y-%m-%d %H:%M}.", "",
             "Errors stop a release. Warnings need a person to look. Notes are for the record.", ""]
    for c in sorted(cohorts, key=lambda c: c.number):
        lines.append(f"## Cohort {c.number} ({c.source}, {len(c.monitors)} monitors)")
        if not c.issues:
            lines.append("- No problems found.")
        for issue in c.issues:
            lines.append(f"- **{issue.level}**: {issue.message}")
        lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path
